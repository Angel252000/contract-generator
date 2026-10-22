"""Exportación del contrato editado (HTML del editor) a un .docx real."""

from html.parser import HTMLParser
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

BLOCKS = {"p", "h1", "h2", "h3", "h4", "li", "div", "blockquote"}


class _DocxBuilder(HTMLParser):
    """Convierte el subconjunto de HTML que produce el editor (títulos, párrafos, listas, negrita, cursiva)."""

    def __init__(self, doc: Document):
        super().__init__(convert_charrefs=True)
        self.doc = doc
        self.paragraph = None
        self.bold = 0
        self.italic = 0
        self.lists: list[str] = []

    def _new_paragraph(self, tag: str):
        if tag == "h1":
            self.paragraph = self.doc.add_heading(level=0)
            self.paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif tag in {"h2", "h3", "h4"}:
            self.paragraph = self.doc.add_heading(level=int(tag[1]) - 1)
        elif tag == "li":
            style = "List Number" if self.lists and self.lists[-1] == "ol" else "List Bullet"
            self.paragraph = self.doc.add_paragraph(style=style)
        else:
            self.paragraph = self.doc.add_paragraph()

    def handle_starttag(self, tag, attrs):
        if tag in BLOCKS:
            self._new_paragraph(tag)
        elif tag in {"ul", "ol"}:
            self.lists.append(tag)
        elif tag in {"strong", "b"}:
            self.bold += 1
        elif tag in {"em", "i"}:
            self.italic += 1
        elif tag == "br" and self.paragraph is not None:
            self.paragraph.add_run().add_break()
        elif tag == "hr":
            self.doc.add_paragraph("_" * 40).alignment = WD_ALIGN_PARAGRAPH.CENTER
            self.paragraph = None

    def handle_endtag(self, tag):
        if tag in BLOCKS:
            self.paragraph = None
        elif tag in {"ul", "ol"} and self.lists:
            self.lists.pop()
        elif tag in {"strong", "b"}:
            self.bold = max(0, self.bold - 1)
        elif tag in {"em", "i"}:
            self.italic = max(0, self.italic - 1)

    def handle_data(self, data):
        if not data.strip() and self.paragraph is None:
            return
        if self.paragraph is None:
            self._new_paragraph("p")
        run = self.paragraph.add_run(data)
        run.bold = bool(self.bold) or None
        run.italic = bool(self.italic) or None


def html_to_docx(html: str) -> bytes:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(11.5)
    builder = _DocxBuilder(doc)
    builder.feed(html)
    builder.close()
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
