import json
from io import BytesIO

from docx import Document

from contract_generator.api.web import create_app
from contract_generator.service import ContractService

from .conftest import SETTINGS, FakeWriter


def sse_events(response):
    return [json.loads(line[6:]) for line in response.get_data(as_text=True).split("\n\n") if line]


def test_index_lists_contract_types(client):
    html = client.get("/").get_data(as_text=True)
    assert html.count('data-type="') == 6
    assert "IA: test-model" in html


def test_healthz(client):
    assert client.get("/healthz").get_json()["status"] == "ok"


def test_generate_streams_events(client, fake_writer):
    res = client.post("/api/generate", json={"description": "Web", "type": "servicios"})
    assert res.mimetype == "text/event-stream"
    events = sse_events(res)
    assert "".join(e.get("text", "") for e in events).startswith("# CONTRATO")
    assert events[-1] == {"done": True}
    assert fake_writer.requests[0].description == "Web"


def test_generate_validates_before_streaming(client):
    res = client.post("/api/generate", json={"description": ""})
    assert res.status_code == 400
    assert "Describe" in res.get_json()["error"]


def test_generate_reports_writer_error_in_stream():
    app = create_app(ContractService(FakeWriter(fail_with="Sin conexión con la API.")), SETTINGS)
    events = sse_events(app.test_client().post("/api/generate", json={"type": "nda"}))
    assert events[-1] == {"error": "Sin conexión con la API."}


def test_edit(client):
    res = client.post("/api/edit", json={"text": "hola", "action": "mejorar"})
    assert res.get_json() == {"text": "[mejorar] hola"}


def test_edit_errors():
    app = create_app(ContractService(FakeWriter(fail_with="Límite")), SETTINGS)
    c = app.test_client()
    assert c.post("/api/edit", json={"text": "hola", "action": "nada"}).status_code == 400
    assert c.post("/api/edit", json={"text": "hola", "action": "acortar"}).status_code == 502


def test_export_docx(client):
    html = "<h1>Contrato de Servicios</h1><h2>PRIMERA. OBJETO</h2><p>Entre <strong>Luna</strong> y otro.</p>"
    res = client.post("/api/export/docx", json={"html": html, "title": "Contrato de Servicios"})
    assert res.status_code == 200
    assert 'filename="contrato-de-servicios.docx"' in res.headers["Content-Disposition"]
    doc = Document(BytesIO(res.data))
    texts = [p.text for p in doc.paragraphs]
    assert texts[:3] == ["Contrato de Servicios", "PRIMERA. OBJETO", "Entre Luna y otro."]
    assert any(run.bold and run.text == "Luna" for run in doc.paragraphs[2].runs)


def test_export_docx_requires_content(client):
    assert client.post("/api/export/docx", json={"html": "  "}).status_code == 400


def test_rejects_oversized_body(client):
    assert client.post("/api/edit", data="x" * 1_100_000, content_type="application/json").status_code == 413
